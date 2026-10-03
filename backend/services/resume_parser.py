import re
import io
from typing import Dict, Any, Tuple
from pypdf import PdfReader
import docx

def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    """Extracts raw text from PDF or DOCX file bytes."""
    ext = filename.split(".")[-1].lower()
    text = ""

    if ext == "pdf":
        try:
            pdf_file = io.BytesIO(file_bytes)
            reader = PdfReader(pdf_file)
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
        except Exception as e:
            text = ""
    elif ext in ["docx", "doc"]:
        try:
            docx_file = io.BytesIO(file_bytes)
            doc = docx.Document(docx_file)
            for para in doc.paragraphs:
                text += para.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    text += " ".join([cell.text for cell in row.cells]) + "\n"
        except Exception as e:
            text = ""

    return text.strip()

def parse_resume_structure(text: str) -> Tuple[Dict[str, Any], Dict[str, str]]:
    """
    Parses resume text into structured components:
    - Contact Information (email, phone, linkedin, github, portfolio, name)
    - Sections (education, experience, projects, skills, certifications, summary)
    """
    contact_info = {
        "email": None,
        "phone": None,
        "linkedin": None,
        "github": None,
        "portfolio": None,
        "name": None
    }

    # Extract Email
    email_match = re.search(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}', text)
    if email_match:
        contact_info["email"] = email_match.group(0)

    # Extract Phone (Robust multi-pattern detection for Indian & International formats)
    # 1. Look for explicit contact/phone labels first
    label_pattern = r'(?:phone|mobile|tel|contact|ph|cell)\s*(?:no\.?|num|number)?\s*[:.-]?\s*(\+?\d[\d\s\(\)\.\-]{7,20}\d)'
    label_match = re.search(label_pattern, text, re.IGNORECASE)
    if label_match:
        raw_phone = label_match.group(1).strip()
        digits = re.sub(r'\D', '', raw_phone)
        if 7 <= len(digits) <= 15:
            contact_info["phone"] = raw_phone

    if not contact_info["phone"]:
        phone_patterns = [
            r'\+\d{1,4}[-.\s]?(?:\(\d{1,5}\)[-.\s]?)?\d{1,5}[-.\s]?\d{2,5}[-.\s]?\d{3,5}',
            r'\(?\d{3,5}\)?[-.\s]?\d{2,5}[-.\s]?\d{3,5}',
            r'\b\d{10,12}\b'
        ]
        for pattern in phone_patterns:
            match = re.search(pattern, text)
            if match:
                candidate = match.group(0).strip()
                digits = re.sub(r'\D', '', candidate)
                if 7 <= len(digits) <= 15:
                    contact_info["phone"] = candidate
                    break

    # Extract LinkedIn / GitHub / Portfolio
    linkedin_match = re.search(r'(https?://)?(www\.)?linkedin\.com/in/[\w-]+', text, re.IGNORECASE)
    if linkedin_match:
        contact_info["linkedin"] = linkedin_match.group(0)

    github_match = re.search(r'(https?://)?(www\.)?github\.com/[\w-]+', text, re.IGNORECASE)
    if github_match:
        contact_info["github"] = github_match.group(0)

    portfolio_match = re.search(r'(https?://)?([\w-]+\.)+(com|io|dev|me|tech)/?', text, re.IGNORECASE)
    if portfolio_match and not linkedin_match and not github_match:
        contact_info["portfolio"] = portfolio_match.group(0)

    # Extract Name (Guessing first non-empty line)
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    if lines:
        contact_info["name"] = lines[0]

    # Section Headers Detection Patterns
    section_patterns = {
        "education": r'\b(education|educational|academic|academics|qualification|qualifications|university|college|school|degree|degrees|schooling)\b',
        "experience": r'\b(experience|experiences|work experience|employment|history|professional experience|work history|career history|internships|internship)\b',
        "projects": r'\b(projects|project|key projects|personal projects|portfolio projects|academic projects)\b',
        "skills": r'\b(skills|skill|technical skills|technologies|expertise|core competencies|competencies|technical expertise)\b',
        "certifications": r'\b(certifications|certification|certificates|certificate|licenses|license|training|courses)\b',
        "summary": r'\b(summary|objective|profile|about me|professional summary|executive summary|overview|career objective)\b'
    }

    sections = {
        "education": "",
        "experience": "",
        "projects": "",
        "skills": "",
        "certifications": "",
        "summary": ""
    }

    seen_sections = set()
    current_section = "summary"
    for line in text.splitlines():
        line_clean = line.strip()
        if not line_clean:
            continue
        
        # Check if line matches a section header
        matched_section = None
        for sec_name, pattern in section_patterns.items():
            if re.search(pattern, line_clean, re.IGNORECASE) and len(line_clean) < 65:
                matched_section = sec_name
                break
        
        if matched_section:
            current_section = matched_section
            seen_sections.add(matched_section)
        else:
            sections[current_section] += line_clean + "\n"

    # Determine sections found
    found_sections_list = []
    for k, v in sections.items():
        if k in seen_sections or len(v.strip()) > 5:
            found_sections_list.append(k)

    # Robust fallback check for Education section
    if "education" not in found_sections_list:
        edu_fallback = r'\b(education|educational|academic|academics|qualification|qualifications|b\.?tech|b\.?e|b\.?sc|bca|mca|m\.?tech|bachelor|bachelors|master|masters|phd|diploma|university|college|schooling|hsc|ssc|cbse|icse)\b'
        if re.search(edu_fallback, text, re.IGNORECASE):
            found_sections_list.append("education")

    structured_data = {
        "contact_info": contact_info,
        "sections_found": list(set(found_sections_list)),
        "word_count": len(text.split()),
        "character_count": len(text)
    }

    return structured_data, sections

